"""
Unit Testing: Python Worker Data Extraction
Verifies multimodal JSON extraction and exact .xlsx formatting with openpyxl.
"""

import os
import openpyxl
import pytest
from backend.worker.data_extractor import DataExtractor


@pytest.fixture
def temp_output_dir(tmp_path):
    out = tmp_path / "spreadsheets"
    out.mkdir()
    return str(out)


def test_data_extractor_formatting(temp_output_dir):
    extractor = DataExtractor(output_dir=temp_output_dir)

    # Mock multimodal LLM extracted JSON output
    mock_extracted_json = [
        {"vendor": "Omni Tech Supply", "item": "NVMe SSD 2TB", "quantity": 2, "unit_price": 149.99, "total_cost": 299.98},
        {"vendor": "Omni Tech Supply", "item": "USB-C Hardware Key", "quantity": 5, "unit_price": 45.00, "total_cost": 225.00},
        {"vendor": "Cloud Infra Inc", "item": "Edge Gateway Module", "quantity": 1, "unit_price": 550.00, "total_cost": 550.00},
    ]

    xlsx_path = extractor.extract_to_excel(
        records=mock_extracted_json,
        filename="test_synthesis.xlsx",
        sheet_title="Hardware_Procurement",
        report_title="Q4 Hardware Ingestion Report",
    )

    assert os.path.exists(xlsx_path)

    # Inspect openpyxl workbook structure and styling
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    ws = wb["Hardware_Procurement"]

    # 1. Title verification
    title_cell = ws.cell(row=1, column=1)
    assert title_cell.value == "Q4 Hardware Ingestion Report"
    assert title_cell.font.bold is True
    assert title_cell.font.color.rgb == "00FFFFFF" or title_cell.font.color.rgb == "FFFFFFFF"

    # 2. Table Headers (Row 3)
    headers = [ws.cell(row=3, column=c).value for c in range(1, 6)]
    assert headers == ["Vendor", "Item", "Quantity", "Unit Price", "Total Cost"]
    assert ws.cell(row=3, column=1).font.bold is True

    # 3. Data Rows formatting
    assert ws.cell(row=4, column=1).value == "Omni Tech Supply"
    assert ws.cell(row=4, column=3).value == 2
    assert ws.cell(row=4, column=4).value == 149.99
    # Currency formatting check
    assert ws.cell(row=4, column=4).number_format == '$#,##0.00'
    assert ws.cell(row=4, column=5).number_format == '$#,##0.00'

    # 4. Total Summary Row (Row 7: header=3, 3 data rows -> total row = 7)
    total_label = ws.cell(row=7, column=1).value
    assert total_label == "Total Summary"

    # Formula check on numeric columns
    total_cost_formula = ws.cell(row=7, column=5).value
    assert total_cost_formula == "=SUM(E4:E6)"
    assert ws.cell(row=7, column=5).number_format == '$#,##0.00'
    assert ws.cell(row=7, column=5).font.bold is True
